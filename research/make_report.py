"""把 overfit.pkl 写成 Excel 报告。python research/make_report.py <overfit.pkl> <out.xlsx>"""
import sys
import pickle

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

R = pickle.load(open(sys.argv[1], "rb"))
OUT = sys.argv[2]
F = "Arial"
HDR = PatternFill("solid", fgColor="305496")
HL = PatternFill("solid", fgColor="FFF2CC")
GOOD = PatternFill("solid", fgColor="E2EFDA"); WARN = PatternFill("solid", fgColor="FCE4D6")
thin = Side(style="thin", color="D9D9D9")

wb = Workbook()


def sheet(title, df, pct_cols=(), money_cols=(), num_cols=(), widths=None, note=None, hl_col=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(name=F, italic=True, color="595959")
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(6, len(df.columns)))
        ws.row_dimensions[1].height = 48
        r0 = 3
    for c, name in enumerate(df.columns, 1):
        x = ws.cell(r0, c, name); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = HDR
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for c, v in enumerate(row, 1):
            if isinstance(v, (np.floating, float)) and not np.isfinite(v):
                v = None
            if isinstance(v, (np.integer,)):
                v = int(v)
            if isinstance(v, (np.floating,)):
                v = float(v)
            if isinstance(v, (np.bool_, bool)):
                v = "是" if v else ""
            if isinstance(v, dict):
                v = "，".join(f"{k}={_fmt(k, x)}" for k, x in v.items())
            x = ws.cell(i, c, v); x.font = Font(name=F); x.border = Border(bottom=thin)
            col = df.columns[c - 1]
            if col in pct_cols:
                x.number_format = "0.0%;-0.0%;0.0%"
            elif col in money_cols:
                x.number_format = "#,##0"
            elif col in num_cols:
                x.number_format = "0.000"
        if hl_col is not None and getattr(row, hl_col, None) in (True, "是"):
            for c in range(1, len(df.columns) + 1):
                ws.cell(i, c).fill = HL
    ws.row_dimensions[r0].height = 36
    for c, name in enumerate(df.columns, 1):
        w = (widths or {}).get(name, max(10, min(40, len(str(name)) * 2 + 2)))
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    return ws, r0


def _fmt(k, x):
    if k == "mcap":
        return f"{x/1e8:.0f}亿"
    if k in ("drop", "tp"):
        return f"{x:.0%}"
    return x


# ---------------------------------------------------------------- 结论
nb = R["邻域"]; wf = R["滚动样本外"]; dist = R["参数组合分布"]; pl = R["安慰剂"]; cost = R["成本压力"]; con = R["集中度"]
wfc = {a: (np.prod(1 + d["样本外_滚动选参收益"]) ** (1 / len(d)) - 1, np.prod(1 + d["样本外_F版固定参数收益"]) ** (1 / len(d)) - 1,
           np.prod(1 + d["当年全部243组中位数"]) ** (1 / len(d)) - 1) for a, d in wf.groupby("账户")}
sig_rng = nb[nb["信号笔数"] > 0]["每笔平均"]
concl = pd.DataFrame([
    dict(检验="① 信号本身是否有效（安慰剂）", 结果=f"主信号每笔平均 {pl.iloc[0]['实际每笔平均']:.2%}，同日随机挑『同样跌5~10%、市值≥80亿』的股票只有 {pl.iloc[0]['随机300次平均']:.2%}；随机300次没有一次超过实际（最好 {pl.iloc[0]['随机最好一次']:.2%}）。补位信号 {pl.iloc[2]['实际每笔平均']:.2%} vs 随机 {pl.iloc[2]['随机300次平均']:.2%}",
         判断="通过", 含义="『首次进前二后回撤』这个选股逻辑本身有超额，不是凭运气；补位信号超额较小"),
    dict(检验="② 参数邻域（信号级）", 结果=f"主信号每笔平均收益在所有参数取值下都为正，范围 {sig_rng.min():.2%} ~ {sig_rng.max():.2%}；前后两半段（2018-21 / 2022-26）也都为正",
         判断="通过", 含义="信号质量是『一片高原』，不是靠某个精确参数撑起来的"),
    dict(检验="③ 参数邻域（账户级）", 结果="账户年化对部分参数很敏感：观察期10→5日 14.6%→5.9%；时间止损3→4日 14.6%→8.0%（但5日又回到14.7%）；进池前2→前5 14.6%→3.4%；跌停家数>3→>5 14.6%→10.4%",
         判断="警示", 含义="账户收益的『形状』受资金占用和路径影响很大，结果非单调 = 有噪音成分；回看天数、跌幅、补仓触发比较稳"),
    dict(检验="④ 选参偏差（243组参数分布）", 结果=f"在 观察期×跌幅×市值×止盈×时间止损 的243组里，F版累计收益排第 {int(dist.iloc[0]['F版排名'])}（稳健）/ 第 {int(dist.iloc[1]['F版排名'])}（进取）；243组年化中位数只有 {dist.iloc[0]['年化中位数']:.1%} / {dist.iloc[1]['年化中位数']:.1%}；约 {dist.iloc[0]['跑赢5万本金的组合占比']:.0%} 的组合能赚钱",
         判断="警示", 含义="F版是从一堆组合里挑出来的头部，14.6%/18.1% 里有『挑选运气』，实际预期应打折"),
    dict(检验="⑤ 滚动样本外（每年只用以前的数据选参）", 结果=f"2020-2026 年化：滚动选参 {wfc['稳健版'][0]:.1%}（稳健）/ {wfc['进取版'][0]:.1%}（进取）；同期F版固定参数 {wfc['稳健版'][1]:.1%} / {wfc['进取版'][1]:.1%}；每年243组中位数 {wfc['稳健版'][2]:.1%} / {wfc['进取版'][2]:.1%}",
         判断="警示", 含义="如果2019年的你按当时最优选参数，之后拿到的大约是10%（稳健）/ 7%（进取），明显低于F版回测。进取版更脆弱"),
    dict(检验="⑥ 交易成本", 结果=f"单边滑点0.1%：稳健年化 {cost.iloc[1]['稳健版_年化']:.1%}；0.2%：{cost.iloc[2]['稳健版_年化']:.1%}；5元最低佣金：{cost.iloc[4]['稳健版_年化']:.1%}；滑点0.2%+5元：{cost.iloc[5]['稳健版_年化']:.1%}（进取 {cost.iloc[5]['进取版_年化']:.1%}）",
         判断="不通过（现实风险最大）", 含义="每笔利润薄（平均+3%左右、止盈只有5%），成本吃掉一大块；补位单金额小，5元最低佣金伤害明显"),
    dict(检验="⑦ 收益集中度", 结果=f"去掉最好的一年（2026）后稳健年化 {con.iloc[0]['去掉最好一年后年化']:.1%}；最赚的20笔贡献了约 {1 - con.iloc[0]['去掉最赚的20笔后总盈亏'] / con.iloc[0]['总盈亏']:.0%} 的总盈利",
         判断="警示", 含义="2026年特别好，拉高了整体；少数大赚的单子很重要"),
    dict(检验="⑧ 逐年稳定性（信号级）", 结果="主信号除2018年（14笔，-0.7%）外每年为正；补位信号9年全部为正但每笔只有0.6%~3.0%",
         判断="通过", 含义="不是靠某一年的行情"),
    dict(检验="⑨ 二次买对照", 结果="每次进池最多买1次 14.6% → 2次 18.9% → 3次 11.6%（稳健），非单调",
         判断="警示", 含义="二次买的提升可能也有噪音成分，继续在看板上并排跑观察，不建议现在切换"),
])
ws, r0 = sheet("结论", concl, widths={"检验": 30, "结果": 90, "判断": 16, "含义": 60},
               note="F版 过拟合检验（2018-01-02 ~ 2026-10-16，5万本金，与《F版_回测报告》同一引擎，默认参数逐笔复现 166,241 / 216,578）。"
                    "结论：选股信号本身是真的（通过安慰剂和邻域检验），但账户层面的 14.6%/18.1% 年化有选参偏差，且对交易成本很敏感。更现实的预期见下方。")
for i in range(len(concl)):
    c = ws.cell(r0 + 1 + i, 3); v = c.value
    c.fill = GOOD if v == "通过" else WARN
    for k in (2, 4):
        ws.cell(r0 + 1 + i, k).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r0 + 1 + i].height = 48
r = r0 + len(concl) + 3
lines = [
    "【更现实的预期】稳健版：无成本时样本外大约 10%~11% 年化；按集合竞价单边 0.1% 滑点、券商5元最低佣金估，大约 6%~9%。进取版样本外更不稳定，不建议作为主力。",
    "【建议1】不要再为了回测好看去微调参数——这类调整正是过拟合的来源。需要改规则时，先在看板上并排跑（像『二次买』那样）。",
    "【建议2】真正的样本外从 2026-10-19 开始：看板每天自动记录。建议攒到约 30 笔主信号（按历史频率约 9~12 个月）再评估：主信号每笔平均 ≥ +1.5%、胜率 ≥ 65% 算与回测一致；每笔 < +0.5% 要重新审视。",
    "【建议3】执行成本是第一风险：务必集合竞价；确认券商是否免5元最低佣金；补位单首笔只有约2000元，若券商有5元最低佣金，可考虑不做补位（关掉补位稳健版年化 13.7%，回撤反而更小 -22.6%）。",
    "【方法局限】所有检验用的是同一段历史；无分钟数据，止盈按日内最高价判断；2026-09-22 后流通市值为推算；幸存者偏差：数据含已退市股票。",
]
for k, t in enumerate(lines):
    c = ws.cell(r + k, 1, t); c.font = Font(name=F, bold=k == 0); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r + k, start_column=1, end_row=r + k, end_column=4); ws.row_dimensions[r + k].height = 36

# ---------------------------------------------------------------- 明细表
nb2 = nb.copy()
nb2["取值"] = [(_fmt({"流通市值门槛": "mcap", "跌幅下限 drop": "drop", "止盈": "tp", "补仓触发": "drop"}.get(p, ""), v)
               if not isinstance(v, str) else v) for p, v in zip(nb2["参数"], nb2["取值"])]
cols = ["参数", "取值", "是否F版", "信号笔数", "每笔平均", "每笔中位", "胜率_信号", "前半段每笔", "后半段每笔",
        "稳健版_期末", "稳健版_年化", "稳健版_最大回撤", "稳健版_最差年", "稳健版_前半段年化", "稳健版_后半段年化",
        "进取版_期末", "进取版_年化", "进取版_最大回撤", "进取版_最差年", "进取版_前半段年化", "进取版_后半段年化"]
nb2 = nb2[cols].rename(columns={"每笔平均": "主信号每笔平均(信号级)", "每笔中位": "主信号每笔中位", "胜率_信号": "主信号胜率",
                                "前半段每笔": "2018-21每笔", "后半段每笔": "2022-26每笔", "信号笔数": "主信号笔数"})
pc = [c for c in nb2.columns if any(k in c for k in ("每笔", "胜率", "年化", "回撤", "最差年"))]
sheet("参数邻域", nb2, pct_cols=pc, money_cols=["稳健版_期末", "进取版_期末"], hl_col="是否F版",
      widths={"参数": 16, "取值": 10},
      note="每次只改一个参数，其余保持F版。黄色行=F版。『信号级』= 每个主信号独立用1万元按止盈/补仓/时间止损模拟，不受资金和持仓影响，最能反映信号本身的质量；『账户级』= 5万本金按F版仓位规则完整模拟。"
           "前半段=2018-2021单独从5万跑，后半段=2022-2026单独从5万跑。")

wf2 = wf.copy()
ws, r0 = sheet("滚动样本外", wf2, pct_cols=["样本外_滚动选参收益", "样本外_F版固定参数收益", "当年全部243组中位数", "当年F版在243组中的分位"],
               widths={"训练期最优参数": 46, "账户": 10, "年份": 8},
               note="每年年初，只用之前年份的数据，在 观察期{5,10,15}×跌幅{4%,5%,6%}×市值{50,80,120亿}×止盈{4%,5%,6%}×时间止损{2,3,4} 共243组里选累计收益最高的参数，用到当年（真正的样本外）。"
                    "对比：同年F版固定参数（F版是用全部历史选出的，含未来信息）。下方为2020-2026连乘年化。")
n = len(wf2); r = r0 + n + 2
ws.cell(r, 1, "2020-2026 连乘年化").font = Font(name=F, bold=True)
for k, (acc, rows_) in enumerate(wf2.groupby("账户", sort=False)):
    first = r0 + 1 + rows_.index[0]; last = r0 + 1 + rows_.index[-1]; m = len(rows_)
    ws.cell(r + 1 + k, 1, acc).font = Font(name=F)
    for col_name in ["样本外_滚动选参收益", "样本外_F版固定参数收益", "当年全部243组中位数"]:
        ci = list(wf2.columns).index(col_name) + 1; L = get_column_letter(ci)
        c = ws.cell(r + 1 + k, ci, f"=SUMPRODUCT(PRODUCT(1+{L}{first}:{L}{last}))^(1/{m})-1")
        c.number_format = "0.0%"; c.font = Font(name=F, bold=True)

sheet("参数组合分布", dist, pct_cols=["跑赢5万本金的组合占比", "年化中位数"], num_cols=["F版累计倍数", "最优累计倍数", "中位数累计倍数", "最差累计倍数"],
      note="同一243组参数在 2018-2026 全样本的累计收益分布。F版排名越靠前、与中位数差距越大，说明『挑出来』的成分越多。")
sheet("安慰剂", pl, pct_cols=["实际每笔平均", "随机300次平均", "随机95分位", "随机最好一次", "实际超过随机的比例"], widths={"对照": 46},
      note="把每个真实信号换成『同一天随机挑的另一只股票』，用完全相同的止盈/补仓/时间止损规则模拟，重复300次。实际结果远高于随机分布 → 选股逻辑有效。")
sheet("成本压力", cost, pct_cols=[c for c in cost.columns if "年化" in c or "回撤" in c or "最差" in c],
      money_cols=[c for c in cost.columns if "期末" in c], widths={"成本设定": 24},
      note="滑点：买入价×(1+滑点)，卖出价×(1-滑点)，开盘、盘中、收盘成交都加。最低佣金：每笔佣金不足5元按5元收。")
sheet("逐年信号", R["逐年信号"], pct_cols=["主信号每笔平均", "主信号胜率", "补位每笔平均", "补位胜率"],
      note="信号级（每个信号独立模拟，不受资金约束），看每年是否都成立。补位信号笔数远多于账户实际成交（账户只在空仓时买补位）。")
sheet("集中度", con, pct_cols=["全部年化", "去掉最好一年后年化"], money_cols=["总盈亏", "去掉最赚的10笔后总盈亏", "去掉最赚的20笔后总盈亏"],
      note="收益是否依赖少数年份/少数交易。")
meth = pd.DataFrame([
    dict(项目="引擎", 说明="research/fast.py：矩阵化重写，默认参数与 engine/core.py、原 final_F.py 逐笔一致（稳健 166,241 / 进取 216,578 / 二次买 228,462 / 208,483）"),
    dict(项目="数据", 说明="谷票主板日线 2018-01-02 ~ 2026-10-16（含已退市股票），与F版报告相同"),
    dict(项目="信号级模拟", 说明="每个信号次日开盘买1万元，按止盈5%/补仓-3%到2倍、2.5倍/时间止损3日/封跌停顺延，独立计算收益率；开盘涨跌停或停牌则不计"),
    dict(项目="账户级模拟", 说明="5万本金，F版仓位、优先级、最多4只、补位仅空仓买等全部规则"),
    dict(项目="前后半段", 说明="账户：2018-2021 与 2022-2026 各自从5万重新开始；信号级：按信号日期分组"),
    dict(项目="滚动样本外", 说明="用每组参数全样本运行的逐年收益近似（年初持仓沿用），每年用之前年份累计收益选最优组合"),
    dict(项目="安慰剂", 说明="随机种子42，每组300次"),
    dict(项目="未检验", 说明="分钟级成交（止盈按日内最高价）、真实集合竞价成交量约束、未来市场制度变化"),
])
sheet("方法说明", meth, widths={"项目": 14, "说明": 120})
del wb["Sheet"]
wb.save(OUT)
print("saved", OUT)
